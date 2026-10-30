import os
import hashlib
import re
from xml.sax.saxutils import escape as xml_escape
import tempfile
import uuid
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Optional, List, Dict, Any, Tuple
from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import openpyxl
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import csv
import io
import json
from pathlib import Path
import shutil
import stat

BUSINESS_RULES = {
    "normal_start_time": "08:00",
    "late_arrival_threshold": "08:30",
    "very_late_arrival_threshold": "09:30",
    "overtime_threshold": "19:30",
    "default_work_end_time": "17:30",
    "weekend_overtime_evaluated_separately": True,
}
EFFECTIVE_RULE_KEYS = ("normal_start_time", "very_late_arrival_threshold", "overtime_threshold")


def effective_rules_version(rules: dict) -> str:
    payload = json.dumps({key: rules[key] for key in EFFECTIVE_RULE_KEYS}, sort_keys=True)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:12]


RULES_VERSION = effective_rules_version(BUSINESS_RULES)


class _Event:
    """Lightweight event object used by the attendance processing logic."""
    __slots__ = ('employee_id', 'employee_name', 'date', 'event_type', 'timestamp', 'source_event_type', 'source_timestamp')

    def __init__(self, employee_id, employee_name, date, event_type, timestamp, source_event_type=None, source_timestamp=None):
        self.employee_id = employee_id
        self.employee_name = employee_name
        self.date = date
        self.event_type = event_type
        self.timestamp = timestamp
        self.source_event_type = source_event_type if source_event_type is not None else event_type
        self.source_timestamp = source_timestamp if source_timestamp is not None else (timestamp.isoformat() if timestamp else None)

app = FastAPI(title="MM&RJD Attendance Report Automation", version="0.1.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

TEMP_DIR = Path(tempfile.gettempdir()) / "mm_attendance"
TEMP_DIR.mkdir(exist_ok=True)


def job_directory(job_id: str, *, create: bool = False) -> Path:
    try:
        normalized_id = str(uuid.UUID(job_id))
    except (ValueError, TypeError, AttributeError):
        raise HTTPException(status_code=400, detail="A valid upload/job ID is required.")
    if TEMP_DIR.is_symlink():
        raise HTTPException(status_code=500, detail="Configured job storage root must not be a symlink.")
    try:
        storage_root = TEMP_DIR.resolve(strict=True)
    except OSError as error:
        raise HTTPException(status_code=500, detail="Configured job storage root is unavailable.") from error
    if not storage_root.is_dir():
        raise HTTPException(status_code=500, detail="Configured job storage root is not a directory.")
    path = storage_root / normalized_id
    if create:
        path.mkdir(parents=True, exist_ok=True)
    if path.is_symlink():
        raise HTTPException(status_code=409, detail="Upload job path is invalid.")
    try:
        resolved_path = path.resolve(strict=True)
    except OSError as error:
        if not create:
            raise HTTPException(status_code=404, detail="Upload job not found. Upload the attendance file again.") from error
        raise HTTPException(status_code=500, detail="Could not create a safe upload job directory.") from error
    if resolved_path.parent != storage_root or not resolved_path.is_dir():
        raise HTTPException(status_code=409, detail="Upload job path is outside the configured storage root.")
    return resolved_path

@app.get("/", summary="Health check")
def health():
    return {"status": "ok", "service": "MM&RJD Attendance Report Automation"}

@app.get("/api/health", summary="API health check")
def api_health():
    return {"status": "ok"}

@app.get("/api/jobs/{job_id}/status", summary="Check whether an upload job has a current saved result")
def get_job_status(job_id: str):
    upload = read_job_json(job_id, "upload.json")
    if not isinstance(upload, dict):
        raise HTTPException(status_code=500, detail="Saved upload metadata is invalid.")
    processed_rules_version = None
    try:
        result = read_job_json(job_id, "processed.json")
        processed = True
    except HTTPException as error:
        if error.status_code != 404:
            raise
        processed = False
    if processed:
        if not isinstance(result, dict):
            raise HTTPException(status_code=500, detail="Saved processing result is unavailable or invalid.")
        if result.get("job_id") != str(uuid.UUID(job_id)):
            raise HTTPException(status_code=409, detail="Saved processing result does not belong to this upload job.")
        version = result.get("rules_version")
        processed_rules_version = version if isinstance(version, str) else None

    return {
        "success": True,
        "job_id": str(uuid.UUID(job_id)),
        "processed": processed,
        "processed_rules_version": processed_rules_version,
        "current_rules_version": effective_rules_version(BUSINESS_RULES),
    }

def cleanup_temp_file(filepath: str):
    try:
        p = Path(filepath)
        if p.exists():
            p.unlink()
    except Exception:
        pass


def read_job_json(job_id: str, filename: str) -> Any:
    """Reject a final-component symlink when opening a job JSON file.

    O_NOFOLLOW protects the leaf file, in addition to job_directory's path
    checks. This path-based open does not eliminate races where a process with
    write access replaces a parent directory between validation and open.
    """
    if filename not in {"upload.json", "processed.json", "roster.json"}:
        raise ValueError("Unsupported job file")
    directory = job_directory(job_id)
    path = directory / filename
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    try:
        fd = os.open(path, flags)
    except FileNotFoundError as error:
        raise HTTPException(status_code=404, detail=f"Job file {filename} was not found.") from error
    except OSError as error:
        raise HTTPException(status_code=409, detail=f"Job file {filename} is not a safe regular file.") from error
    try:
        if not stat.S_ISREG(os.fstat(fd).st_mode):
            raise HTTPException(status_code=409, detail=f"Job file {filename} is not a regular file.")
        with os.fdopen(fd, "r", encoding="utf-8") as stream:
            fd = -1
            return json.load(stream)
    finally:
        if fd >= 0:
            os.close(fd)


async def read_upload_bounded(file: UploadFile, max_size: int, detail: str) -> bytes:
    """Read at most max_size plus one byte, enough to reject oversized input."""
    try:
        content = await file.read(max_size + 1)
    except Exception as error:
        await file.close()
        raise HTTPException(status_code=400, detail="Upload was interrupted while receiving the file.") from error
    if len(content) > max_size:
        await file.close()
        raise HTTPException(status_code=413, detail=detail)
    return content

def detect_ivms_format(file_content: bytes, filename: str) -> Dict[str, Any]:
    ext = Path(filename).suffix.lower()
    if ext == '.csv':
        return parse_ivms_csv(file_content)
    elif ext in ('.xlsx', '.xls'):
        return parse_ivms_excel(file_content)
    elif ext == '.pdf':
        return parse_ivms_pdf(file_content)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format: {ext}")

def parse_ivms_csv(file_content: bytes) -> Dict[str, Any]:
    text = file_content.decode('utf-8', errors='replace')
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise HTTPException(status_code=400, detail="Empty CSV file")
    headers = [h.strip().lower() for h in rows[0]]
    data_rows = [row for row in rows[1:] if any(cell.strip() for cell in row)]
    column_map = {}
    for i, header in enumerate(headers):
        if any(kw in header for kw in ('employee name', 'person name', 'staff name', 'name')):
            column_map['employee_name'] = i
        elif any(kw in header for kw in ('employee id', 'person id', 'staff id', 'card no', 'card number', 'employee no', 'employee number', 'id')):
            column_map['employee_id'] = i
        elif any(kw in header for kw in ('event type', 'event', 'type', 'action', 'direction')):
            column_map['event_type'] = i
        elif any(kw in header for kw in ('timestamp', 'date time', 'datetime', 'time')):
            column_map['timestamp'] = i
        elif any(kw in header for kw in ['date']):
            column_map['date'] = i
        elif any(kw in header for kw in ['device']):
            column_map['device_id'] = i

    # Preserve the legacy four-column export layout previously supported by
    # positional fallbacks: employee ID, name/other, event type, timestamp.
    # Require a plausible width before using it so unrelated CSVs cannot turn
    # arbitrary values into attendance events.
    if not {'employee_id', 'event_type', 'timestamp'}.issubset(column_map):
        positional_layout = (
            len(headers) >= 4
            and any(token in headers[0] for token in ('employee', 'person', 'staff', 'worker', 'card', ' id'))
            and any(token in headers[2] for token in ('event', 'type', 'action', 'direction'))
            and any(token in headers[3] for token in ('time', 'timestamp', 'datetime'))
        )
        if positional_layout:
            column_map.setdefault('employee_id', 0)
            column_map.setdefault('event_type', 2)
            column_map.setdefault('timestamp', 3)
            if len(headers) >= 5:
                column_map.setdefault('date', 4)
            if len(headers) >= 2:
                column_map.setdefault('employee_name', 1)
    if not {'employee_id', 'event_type', 'timestamp'}.issubset(column_map):
        raise HTTPException(
            status_code=400,
            detail="CSV headers must identify employee ID, event type, and timestamp columns, or use the supported four-column positional export format.",
        )
    if not data_rows:
        raise HTTPException(status_code=400, detail="CSV contains a header but no attendance rows.")

    for row_number, row in enumerate(data_rows, start=2):
        required_indexes = (column_map['employee_id'], column_map['event_type'], column_map['timestamp'])
        if any(index >= len(row) or not str(row[index]).strip() for index in required_indexes):
            raise HTTPException(status_code=400, detail=f"CSV row {row_number} is missing employee ID, event type, or timestamp.")
        try:
            timestamp = parse_timestamp(str(row[column_map['timestamp']]))
        except (ValueError, IndexError):
            raise HTTPException(status_code=400, detail=f"CSV row {row_number} has an invalid or unsupported timestamp.")
        if 'date' in column_map:
            if column_map['date'] >= len(row) or not str(row[column_map['date']]).strip():
                raise HTTPException(status_code=400, detail=f"CSV row {row_number} is missing its date value.")
            try:
                date_value = parse_timestamp(str(row[column_map['date']])).date()
            except (ValueError, IndexError):
                raise HTTPException(status_code=400, detail=f"CSV row {row_number} has an invalid or unsupported date.")
            if date_value != timestamp.date():
                raise HTTPException(status_code=400, detail=f"CSV row {row_number} date does not match its timestamp date.")
    return {'format': 'csv', 'column_map': column_map, 'headers': headers,
            'total_rows': len(data_rows), 'raw_data': data_rows}

def parse_ivms_excel(file_content: bytes) -> Dict[str, Any]:
    import openpyxl as _pxl
    wb = _pxl.load_workbook(io.BytesIO(file_content), data_only=True)
    sheet = wb.active
    headers = [cell.value for cell in sheet[1]] if sheet[1] else []
    data_rows = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        data_rows.append([cell for cell in row])
    if not headers or all(h is None for h in headers):
        raise HTTPException(status_code=400, detail="Excel file has no header row")
    headers = [str(h).strip().lower() if h else '' for h in headers]
    column_map = {}
    for i, header in enumerate(headers):
        if 'name' in header:
            column_map['employee_name'] = i
        elif any(kw in header for kw in ['employee id', 'id', 'card', 'staff']):
            column_map['employee_id'] = i
        elif any(kw in header for kw in ['date']):
            column_map['date'] = i
        elif any(kw in header for kw in ['event', 'type', 'action']):
            column_map['event_type'] = i
        elif any(kw in header for kw in ['time', 'timestamp']):
            column_map['timestamp'] = i
        elif any(kw in header for kw in ['device']):
            column_map['device_id'] = i
    if 'timestamp' not in column_map and len(headers) > 3:
        column_map['timestamp'] = 3
    if 'event_type' not in column_map and len(headers) > 2:
        column_map['event_type'] = 2
    if 'employee_id' not in column_map:
        column_map['employee_id'] = 0
    if 'employee_name' not in column_map and len(headers) > 1:
        column_map['employee_name'] = 1
    if 'date' not in column_map and len(headers) > 4:
        column_map['date'] = 4
    return {'format': 'excel', 'column_map': column_map, 'headers': headers,
        'total_rows': len(data_rows), 'raw_data': data_rows}

def parse_ivms_pdf(file_content: bytes) -> Dict[str, Any]:
    try:
        import pdfminer.high_level as pdfml
        text = pdfml.extract_text(io.BytesIO(file_content))
    except ImportError:
        raise HTTPException(status_code=400, detail="pdfminer not available. Use Excel/CSV export from iVMS-4200.")
    if not text:
        raise HTTPException(status_code=400, detail="Could not extract text from PDF.")
    lines = text.split('\n')
    data_rows = []
    headers = []
    column_map = {}
    for i, line in enumerate(lines):
        line = line.strip()
        if not line:
            continue
        parts = line.split()
        if len(parts) >= 3:
            if i == 0:
                headers = [p.strip().lower() for p in parts]
                for j, h in enumerate(headers):
                    if any(kw in h for kw in ['employee', 'id', 'card']):
                        column_map['employee_id'] = j
                    elif any(kw in h for kw in ['name']):
                        column_map['employee_name'] = j
                    elif any(kw in h for kw in ['date']):
                        column_map['date'] = j
                    elif any(kw in h for kw in ['time', 'timestamp']):
                        column_map['timestamp'] = j
                    elif any(kw in h for kw in ['event', 'type']):
                        column_map['event_type'] = j
            else:
                data_rows.append(parts)
    if not column_map:
        raise HTTPException(status_code=400, detail="Could not detect column structure in PDF.")
    return {'format': 'pdf', 'column_map': column_map, 'headers': headers,
        'total_rows': len(data_rows), 'raw_data': data_rows}

def parse_timestamp(ts_str: str) -> datetime:
    for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%d/%m/%Y %H:%M:%S', '%m/%d/%Y %H:%M:%S', '%Y-%m-%d']:
        try:
            return datetime.strptime(ts_str.strip(), fmt)
        except ValueError:
            continue
    raise ValueError(f"Could not parse timestamp: {ts_str}")

def is_weekend(date_str: str) -> bool:
    date = datetime.strptime(date_str, '%Y-%m-%d')
    return date.weekday() >= 5

def process_attendance_events(events: List, rules: dict) -> dict:
    def direction(event: Any) -> str:
        kind = str(event.event_type or "").strip().lower()
        if kind in ("check_in", "check-in", "in"):
            return "in"
        if kind in ("check_out", "check-out", "out"):
            return "out"
        if kind in ("break_start", "break start", "break beginn", "break_end", "break end"):
            return "auxiliary"
        return "other"

    sorted_events = sorted(events, key=lambda e: e.timestamp or datetime.min)
    check_ins = [e for e in sorted_events if direction(e) == "in" and e.timestamp]
    check_outs = [e for e in sorted_events if direction(e) == "out" and e.timestamp]
    unknown_events = [e for e in sorted_events if direction(e) == "other"]
    missing_timestamps = [e for e in sorted_events if e.timestamp is None]
    timestamp_counts: Dict[datetime, int] = {}
    for event in sorted_events:
        if event.timestamp:
            timestamp_counts[event.timestamp] = timestamp_counts.get(event.timestamp, 0) + 1
    duplicate_timestamps = any(count > 1 for count in timestamp_counts.values())

    # A punch is confirmed only when the source explicitly identifies its direction.
    # Repeated/missing/unknown evidence makes the whole day provisional for review.
    check_in = min((e.timestamp for e in check_ins), default=None)
    check_out = max((e.timestamp for e in check_outs), default=None)
    break_starts = [e for e in sorted_events if str(e.event_type).lower() in ("break_start", "break start", "break beginn") and e.timestamp]
    break_ends = [e for e in sorted_events if str(e.event_type).lower() in ("break_end", "break end") and e.timestamp]
    break_duration_minutes = 0
    if break_starts and break_ends and break_ends[0].timestamp >= break_starts[0].timestamp:
        break_duration_minutes = int((break_ends[0].timestamp - break_starts[0].timestamp).total_seconds() / 60)

    late_minutes = 0
    very_late = False
    if check_in:
        event_date = check_in.date()
        ns = datetime.combine(event_date, time.fromisoformat(rules['normal_start_time']))
        late_minutes = max(0, int((check_in - ns).total_seconds() / 60))
        vl = datetime.combine(event_date, time.fromisoformat(rules['very_late_arrival_threshold']))
        very_late = check_in >= vl
    overtime_minutes = 0
    weekend_policy_pending = bool(check_out and is_weekend(check_out.strftime('%Y-%m-%d')))
    if check_out and not weekend_policy_pending:
        ot = datetime.combine(check_out.date(), time.fromisoformat(rules['overtime_threshold']))
        overtime_minutes = max(0, int((check_out - ot).total_seconds() / 60))

    issues = []
    review_notes = []
    has_events = bool(sorted_events)
    if not has_events:
        status = "no_events"
        issues.append("No attendance events recorded. This does not confirm absence.")
    else:
        if unknown_events:
            issues.append("Unknown event type(s) retained as source evidence; direction requires human review.")
        if missing_timestamps:
            issues.append("One or more source events have missing or invalid timestamps.")
        if duplicate_timestamps:
            issues.append("Duplicate event timestamps make punch direction/order ambiguous.")
        if len(check_ins) != 1 or len(check_outs) != 1:
            issues.append("A single explicit check-in and check-out could not be established.")
        if check_in and check_out and check_out <= check_in:
            issues.append("Check-out is not later than check-in.")
        if weekend_policy_pending:
            overtime_minutes = 0
            issues.append("Weekend overtime policy is unresolved; overtime is not calculated for this day.")
        if late_minutes > 60:
            review_notes.append(f"Arrived {late_minutes} minutes after normal start time ({rules['normal_start_time']}); review required.")
        elif very_late:
            review_notes.append(f"Arrival met the very-late threshold ({rules['very_late_arrival_threshold']}); review required.")
        if check_in is None or check_out is None:
            issues.append("Missing punches are not interpreted as confirmed absence.")
        status = "complete" if not issues else ("incomplete" if check_in is None or check_out is None else "ambiguous")
    flagged_for_review = bool(issues)
    status = "complete" if not issues else status
    flagged_for_review = flagged_for_review or bool(review_notes)
    if status != "complete":
        # Do not produce decision-support metrics from a day marked uncertain.
        late_minutes = 0
        very_late = False
        overtime_minutes = 0
    notes = " ".join(issues + review_notes)
    return {
        'employee_id': sorted_events[0].employee_id if sorted_events else '',
        'employee_name': sorted_events[0].employee_name if sorted_events else '',
        'date': sorted_events[0].date if sorted_events else '',
        'events': sorted_events,
        'event_evidence': [
            {'event_type': str(e.source_event_type or ''), 'timestamp': e.source_timestamp, 'timestamp_valid': e.timestamp is not None}
            for e in sorted_events
        ],
        'check_in': check_in,
        'check_out': check_out,
        'break_duration_minutes': break_duration_minutes,
        'late_minutes': late_minutes,
        'very_late': very_late,
        'overtime_minutes': overtime_minutes,
        'status': status,
        'flagged_for_review': flagged_for_review,
        'punches_provisional': status != 'complete',
        'notes': notes
    }

def process_employee_attendance(events: List, rules: dict, roster: dict = None) -> List[dict]:
    events_by_emp_date: dict = {}
    for event in events:
        key = (event.employee_id, event.date)
        if key not in events_by_emp_date:
            events_by_emp_date[key] = []
        events_by_emp_date[key].append(event)
    processed_days = []
    for (emp_id, date_str), day_events in events_by_emp_date.items():
        day_processed = process_attendance_events(day_events, rules)
        processed_days.append(day_processed)
    if roster:
        event_keys = set(events_by_emp_date.keys())
        for emp_id, emp_data in roster.items():
            dates = []
            if isinstance(emp_data, dict) and 'dates' in emp_data:
                dates = emp_data['dates']
            elif isinstance(emp_data, dict) and 'shifts' in emp_data:
                dates = [s['date'] for s in emp_data.get('shifts', []) if isinstance(s, dict) and 'date' in s]
            for date_str in dates:
                if (emp_id, date_str) not in event_keys:
                    emp_name = emp_data.get('name', '') if isinstance(emp_data, dict) else ''
                    day_processed = {
                        'employee_id': emp_id,
                        'employee_name': emp_name,
                        'date': date_str,
                        'events': [],
                        'event_evidence': [],
                        'check_in': None,
                        'check_out': None,
                        'break_duration_minutes': 0,
                        'late_minutes': 0,
                        'very_late': False,
                        'overtime_minutes': 0,
                        'status': 'no_events',
                        'flagged_for_review': True,
                        'punches_provisional': True,
                        'notes': "No attendance events in uploaded file. This does not confirm absence; verify against source records."
                    }
                    processed_days.append(day_processed)
    processed_days.sort(key=lambda d: d['date'])
    return processed_days

# --- Excel Report Generation ---
def generate_excel_report(processed_days: List[dict], rules: dict, metadata: dict = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    metadata_sheet = wb.create_sheet("Report Metadata")
    metadata_sheet.append(["Report source", (metadata or {}).get("source_filename", "Unknown")])
    metadata_sheet.append(["Upload job ID", (metadata or {}).get("job_id", "Unknown")])
    metadata_sheet.append(["Effective rules version", (metadata or {}).get("rules_version", "Unknown")])
    metadata_sheet.append(["Normal start", rules.get("normal_start_time", "")])
    metadata_sheet.append(["Very late threshold", rules.get("very_late_arrival_threshold", "")])
    metadata_sheet.append(["Overtime threshold", rules.get("overtime_threshold", "")])
    metadata_sheet.column_dimensions["A"].width = 28
    metadata_sheet.column_dimensions["B"].width = 50
    header_font = openpyxl.styles.Font(bold=True, color="FFFFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = openpyxl.styles.Border(
        left=openpyxl.styles.Side(style='thin'),
        right=openpyxl.styles.Side(style='thin'),
        top=openpyxl.styles.Side(style='thin'),
        bottom=openpyxl.styles.Side(style='thin')
    )
    headers = [
        "Employee ID", "Employee Name", "Date", "Check-In", "Check-Out",
        "Break (min)", "Late (min)", "Very Late", "Overtime (min)",
        "Status", "Flagged for Review", "Source event evidence", "Notes"
    ]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
    for row_idx, day in enumerate(processed_days, 2):
        check_in_str = day['check_in'].strftime('%H:%M') if day['check_in'] else ""
        check_out_str = day['check_out'].strftime('%H:%M') if day['check_out'] else ""
        evidence = "; ".join(
            f"{event.get('event_type', 'unknown')}: {event.get('timestamp') or 'missing timestamp'}"
            + (" (invalid timestamp)" if event.get('timestamp_valid') is False else "")
            for event in day.get('event_evidence', [])
        )
        if day.get('punches_provisional', day['flagged_for_review']):
            check_in_str = f"Provisional: {check_in_str}" if check_in_str else "Unknown"
            check_out_str = f"Provisional: {check_out_str}" if check_out_str else "Unknown"
        row_data = [
            day['employee_id'],
            day['employee_name'],
            day['date'],
            check_in_str,
            check_out_str,
            day['break_duration_minutes'],
            day['late_minutes'],
            "Yes" if day['very_late'] else "No",
            day['overtime_minutes'],
            day['status'],
            "Yes" if day['flagged_for_review'] else "No",
            evidence,
            (day['notes'][:50] + "...") if len(day['notes']) > 50 else day['notes']
        ]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
        if day.get('punches_provisional', day['flagged_for_review']):
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = openpyxl.styles.PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    column_widths = [15, 20, 12, 18, 18, 12, 12, 12, 15, 15, 15, 45, 50]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

# --- PDF Report Generation ---
def generate_pdf_report(processed_days: List[dict], rules: dict, metadata: dict = None) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), pageCompression=0)
    styles = getSampleStyleSheet()
    header_cell_style = ParagraphStyle(
        "AttendanceHeaderCell", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=6, leading=7, textColor=colors.white, alignment=1,
    )
    body_cell_style = ParagraphStyle(
        "AttendanceBodyCell", parent=styles["Normal"], fontName="Helvetica",
        fontSize=6, leading=7, alignment=1,
    )
    elements = []
    title_style = styles['Title']
    elements.append(Paragraph("MM&RJD Attendance Report", title_style))
    elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
    if metadata:
        elements.append(Paragraph(f"Source file: {xml_escape(str(metadata.get('source_filename', 'Unknown')))}", styles['Normal']))
        elements.append(Paragraph(f"Upload job ID: {xml_escape(str(metadata.get('job_id', 'Unknown')))}", styles['Normal']))
        elements.append(Paragraph(f"Effective rules version: {xml_escape(str(metadata.get('rules_version', 'Unknown')))}", styles['Normal']))
    elements.append(Paragraph("", styles['Normal']))
    elements.append(Paragraph("Attendance Summary", styles['Heading2']))
    elements.append(Paragraph("", styles['Normal']))
    table_data = [[
        "Employee ID", "Employee Name", "Date", "Check-In", "Check-Out",
        "Break (min)", "Late (min)", "Very Late", "Overtime (min)",
        "Status", "Flagged", "Source event evidence", "Notes"
    ]]
    for day in processed_days:
        check_in_str = day['check_in'].strftime('%H:%M') if day['check_in'] else ""
        check_out_str = day['check_out'].strftime('%H:%M') if day['check_out'] else ""
        very_late_str = "Yes" if day['very_late'] else "No"
        flagged_str = "Yes" if day['flagged_for_review'] else "No"
        evidence = "; ".join(
            f"{event.get('event_type', 'unknown')}: {event.get('timestamp') or 'missing timestamp'}"
            + (" (invalid timestamp)" if event.get('timestamp_valid') is False else "")
            for event in day.get('event_evidence', [])
        )
        if day['flagged_for_review']:
            check_in_str = f"Provisional: {check_in_str}" if check_in_str else "Unknown"
            check_out_str = f"Provisional: {check_out_str}" if check_out_str else "Unknown"
        table_data.append([
            day['employee_id'],
            day['employee_name'],
            day['date'],
            check_in_str,
            check_out_str,
            str(day['break_duration_minutes']),
            str(day['late_minutes']),
            very_late_str,
            str(day['overtime_minutes']),
            day['status'],
            flagged_str,
            evidence,
            (day['notes'][:30] + "...") if len(day['notes']) > 30 else day['notes']
        ])
    wrapped_table_data = [
        [
            Paragraph(xml_escape(str(value)).replace("\n", "<br/>"), header_cell_style if row_index == 0 else body_cell_style)
            for value in row
        ]
        for row_index, row in enumerate(table_data)
    ]
    table = Table(wrapped_table_data, repeatRows=1, colWidths=[0.6*72, 0.95*72, 0.75*72, 0.7*72, 0.7*72,
        0.45*72, 0.45*72, 0.45*72, 0.55*72, 0.55*72, 0.6*72, 1.55*72, 1.4*72])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 0), (-1, 0), 12),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
    ]))
    elements.append(table)
    doc.build(elements)
    output.seek(0)
    return output.read()

# --- API Endpoints ---

@app.post("/api/upload", summary="Upload iVMS attendance export file")
async def upload_file(file: UploadFile = File(...)):
    max_size = 50 * 1024 * 1024
    content = await read_upload_bounded(file, max_size, "File too large. Maximum size is 50 MB.")
    filename = file.filename or "uploaded_file"
    ext = Path(filename).suffix.lower()
    if ext not in ('.csv', '.xlsx', '.xls', '.pdf'):
        raise HTTPException(status_code=400, detail=f"Unsupported file format: {ext}. Supported formats: CSV, Excel (.xlsx, .xls), PDF")
    temp_id = str(uuid.uuid4())
    job_dir = job_directory(temp_id, create=True)
    source_path = job_dir / f"source{ext}"
    try:
        with open(source_path, 'wb') as f:
            f.write(content)
        parsed = detect_ivms_format(content, filename)
        events = extract_attendance_events(parsed)
        results_path = job_dir / "upload.json"
        events_data = [(e[0], e[1], e[2], e[3], e[4].isoformat() if e[4] else None, e[5], e[6]) for e in events]
        results = {'employee_events': events_data, 'column_map': parsed['column_map'], 'format': parsed['format'], 'filename': filename}
        with open(results_path, 'w') as f:
            json.dump(results, f)
        return {
            'success': True,
            'job_id': temp_id,
            'filename': filename,
            'format': parsed['format'],
            'column_map': parsed['column_map'],
            'total_records': parsed['total_rows'],
            'events_found': len(events),
            'message': f"File '{filename}' parsed successfully. {len(events)} attendance events found."
        }
    except HTTPException:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise
    except Exception as e:
        shutil.rmtree(job_dir, ignore_errors=True)
        raise HTTPException(status_code=500, detail=f"Error processing file: {str(e)}")
    finally:
        await file.close()

@app.post("/api/roster", summary="Upload employee roster (CSV/Excel)")
async def upload_roster(job_id: str = Form(...), file: UploadFile = File(...)):
    job_dir = job_directory(job_id)
    read_job_json(job_id, "upload.json")
    max_size = 10 * 1024 * 1024
    content = await read_upload_bounded(file, max_size, "Roster file too large. Maximum size is 10 MB.")
    filename = file.filename or "roster_file"
    ext = Path(filename).suffix.lower()
    if ext not in ('.csv', '.xlsx'):
        raise HTTPException(status_code=400, detail="Unsupported roster format. Use CSV or .xlsx Excel.")
    try:
        if ext == '.csv':
            text = content.decode('utf-8', errors='replace')
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
            if not rows:
                raise HTTPException(status_code=400, detail="Empty CSV file")
            headers = [h.strip().lower() for h in rows[0]]
            roster = {}
            for row in rows[1:]:
                if not row or all(not c for c in row):
                    continue
                emp_id = str(row[0]).strip() if len(row) > 0 else ""
                emp_name = str(row[1]).strip() if len(row) > 1 else ""
                dates = [str(c).strip() for c in row[2:]] if len(row) > 2 else []
                if emp_id:
                    roster[emp_id] = {'name': emp_name, 'dates': dates}
        elif ext == '.xlsx':
            import openpyxl as _pxl
            wb = _pxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(status_code=400, detail="Empty Excel file")
            roster = {}
            for row in rows[1:]:
                if not row or all(not c for c in row):
                    continue
                emp_id = str(row[0]).strip() if row[0] is not None else ""
                emp_name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
                dates = [str(c).strip() for c in row[2:]] if len(row) > 2 else []
                if emp_id:
                    roster[emp_id] = {'name': emp_name, 'dates': dates}
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported roster format: {ext}. Use CSV or Excel.")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error reading roster: {str(e)}")

    roster_path = job_dir / "roster.json"
    roster_temp_path = job_dir / f"roster-{uuid.uuid4().hex}.tmp"
    try:
        roster_temp_path.write_text(json.dumps(roster), encoding="utf-8")
        roster_temp_path.replace(roster_path)
    except Exception as error:
        cleanup_temp_file(str(roster_temp_path))
        raise HTTPException(status_code=500, detail="Could not save the roster for this upload job.") from error
    finally:
        await file.close()

    return {
        'success': True,
        'job_id': str(uuid.UUID(job_id)),
        'filename': filename,
        'employees': len(roster),
        'message': f"Roster '{filename}' loaded with {len(roster)} employees."
    }

def extract_attendance_events(parsed: dict) -> List[tuple]:
    events = []
    data = parsed.get('raw_data', [])
    column_map = parsed.get('column_map', {})
    if not data:
        return events
    num_cols = max(column_map.values()) + 1 if column_map else len(data[0]) if data else 0
    for row in data:
        try:
            employee_id = ""
            employee_name = ""
            date_str = ""
            event_type = ""
            raw_event_type = ""
            timestamp_str = ""
            if 'employee_id' in column_map and column_map['employee_id'] < num_cols:
                employee_id = str(row[column_map['employee_id']]).strip() if column_map['employee_id'] < len(row) else ""
            if 'employee_name' in column_map and column_map['employee_name'] < num_cols:
                employee_name = str(row[column_map['employee_name']]).strip() if column_map['employee_name'] < len(row) else ""
            if 'date' in column_map and column_map['date'] < num_cols:
                date_str = str(row[column_map['date']]).strip() if column_map['date'] < len(row) else ""
            if 'event_type' in column_map and column_map['event_type'] < num_cols:
                raw_value = row[column_map['event_type']] if column_map['event_type'] < len(row) else None
                raw_event_type = str(raw_value).strip() if raw_value is not None else ""
                event_type = raw_event_type.lower()
            if 'timestamp' in column_map and column_map['timestamp'] < num_cols:
                raw_value = row[column_map['timestamp']] if column_map['timestamp'] < len(row) else None
                timestamp_str = str(raw_value).strip() if raw_value is not None else ""
            try:
                date_obj = parse_timestamp(date_str) if date_str else None
                date_str = date_obj.strftime('%Y-%m-%d') if date_obj else ""
            except (ValueError, IndexError):
                date_str = ""
            try:
                ts = parse_timestamp(timestamp_str) if timestamp_str else None
            except ValueError:
                ts = None
            et = event_type.lower()
            if et in ('in', 'check-in', 'check_in', 'arrival', 'checkin'):
                et = 'check_in'
            elif et in ('out', 'check-out', 'check_out', 'departure', 'checkout'):
                et = 'check_out'
            elif et in ('break start', 'break_start', 'break beginn'):
                et = 'break_start'
            elif et in ('break end', 'break_end', 'break end'):
                et = 'break_end'
            else:
                et = raw_event_type  # retain source spelling for review evidence
            events.append((employee_id, employee_name, date_str, et, ts, raw_event_type, timestamp_str))
        except Exception:
            continue
    return events

class ProcessRequest(BaseModel):
    job_id: str


@app.post("/api/process", summary="Process attendance events for an explicit upload job")
def process_attendance_endpoint(request: ProcessRequest):
    global RULES_VERSION
    job_dir = job_directory(request.job_id)
    try:
        results = read_job_json(request.job_id, "upload.json")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error reading processed data: {str(e)}")
    employee_events = results.get('employee_events', [])
    events = []
    for event_row in employee_events:
        emp_id, emp_name, date_str, event_type, ts_str = event_row[:5]
        source_event_type = event_row[5] if len(event_row) > 5 else event_type
        source_timestamp = event_row[6] if len(event_row) > 6 else ts_str
        try:
            ts = datetime.fromisoformat(ts_str) if ts_str else None
        except (ValueError, TypeError):
            ts = None
        events.append(_Event(emp_id, emp_name, date_str, event_type, ts, source_event_type, source_timestamp))
    roster = None
    roster_path = job_dir / "roster.json"
    if roster_path.exists():
        try:
            roster = read_job_json(request.job_id, "roster.json")
        except Exception as e:
            raise HTTPException(status_code=500, detail="Saved employee roster is unavailable or invalid.") from e
    rules = dict(BUSINESS_RULES)
    RULES_VERSION = effective_rules_version(rules)
    processed_days = process_employee_attendance(events, rules, roster)
    total_days = len(processed_days)
    present_days = sum(1 for d in processed_days if d['status'] == 'complete')
    late_count = sum(1 for d in processed_days if d['late_minutes'] > 0 and d['status'] in ('complete', 'incomplete'))
    very_late_count = sum(1 for d in processed_days if d['very_late'])
    total_overtime = sum(d['overtime_minutes'] for d in processed_days)
    result = {
        'success': True,
        'job_id': str(uuid.UUID(request.job_id)),
        'source_filename': results.get('filename', ''),
        'rules_version': RULES_VERSION,
        'rules': {key: rules[key] for key in EFFECTIVE_RULE_KEYS},
        'summary': {
            'total_days': total_days,
            'present_days': present_days,
            'incomplete_or_no_events': total_days - present_days,
            'late_count': late_count,
            'very_late_count': very_late_count,
            'total_overtime_minutes': total_overtime,
        },
        'processed_days': [{
            'employee_id': d['employee_id'],
            'employee_name': d['employee_name'],
            'date': d['date'],
            'check_in': d['check_in'].strftime('%H:%M') if d['check_in'] else "",
            'check_out': d['check_out'].strftime('%H:%M') if d['check_out'] else "",
            'break_duration_minutes': d['break_duration_minutes'],
            'late_minutes': d['late_minutes'],
            'very_late': d['very_late'],
            'overtime_minutes': d['overtime_minutes'],
            'status': d['status'],
            'flagged_for_review': d['flagged_for_review'],
            'punches_provisional': d['punches_provisional'],
            'notes': d['notes'],
            'event_evidence': d['event_evidence']
        } for d in processed_days]
    }
    result_path = job_dir / "processed.json"
    temp_result_path = job_dir / f"processed-{uuid.uuid4().hex}.tmp"
    temp_result_path.write_text(json.dumps(result), encoding="utf-8")
    temp_result_path.replace(result_path)
    return result

import base64

class ExportRequest(BaseModel):
    job_id: str
    rules_version: str


def processed_days_for_export(request: ExportRequest) -> Tuple[List[dict], dict, dict]:
    job_dir = job_directory(request.job_id)
    result_path = job_dir / "processed.json"
    if not result_path.is_file():
        raise HTTPException(status_code=404, detail="This upload job has not been processed yet.")
    result = read_job_json(request.job_id, "processed.json")
    if result.get("job_id") != str(uuid.UUID(request.job_id)):
        raise HTTPException(status_code=409, detail="Processed result does not belong to this upload job.")
    if result.get("rules_version") != request.rules_version:
        raise HTTPException(status_code=409, detail="Rules version does not match the processed result. Reprocess this upload before exporting.")
    if result.get("rules_version") != effective_rules_version(BUSINESS_RULES):
        raise HTTPException(status_code=409, detail="Business rules changed after processing. Reprocess this upload before exporting.")
    rules = result.get("rules") or {}
    processed_days = [
        {
            'employee_id': d['employee_id'],
            'employee_name': d['employee_name'],
            'date': d['date'],
            'check_in': datetime.strptime(d['check_in'], '%H:%M') if d['check_in'] else None,
            'check_out': datetime.strptime(d['check_out'], '%H:%M') if d['check_out'] else None,
            'break_duration_minutes': d['break_duration_minutes'],
            'late_minutes': d['late_minutes'],
            'very_late': d['very_late'],
            'overtime_minutes': d['overtime_minutes'],
            'status': d['status'],
            'flagged_for_review': d['flagged_for_review'],
            'punches_provisional': d.get('punches_provisional', d['flagged_for_review']),
            'notes': d['notes'],
            'event_evidence': d.get('event_evidence', [])
        }
        for d in result.get('processed_days', [])
    ]
    metadata = {
        "job_id": result["job_id"],
        "source_filename": result.get("source_filename", "Unknown"),
        "rules_version": result["rules_version"],
    }
    return processed_days, rules, metadata


@app.post("/api/export/excel", summary="Generate Excel attendance report for a processed job")
def export_excel_endpoint(request: ExportRequest):
    processed_days, rules, metadata = processed_days_for_export(request)
    excel_bytes = generate_excel_report(processed_days, rules, metadata)
    return {"success": True, "filename": "attendance_report.xlsx", "job_id": request.job_id, "rules_version": request.rules_version, "data": base64.b64encode(excel_bytes).decode('utf-8')}

@app.post("/api/export/pdf", summary="Generate PDF attendance report")
def export_pdf_endpoint(request: ExportRequest):
    processed_days, rules, metadata = processed_days_for_export(request)
    pdf_bytes = generate_pdf_report(processed_days, rules, metadata)
    return {"success": True, "filename": "attendance_report.pdf", "job_id": request.job_id, "rules_version": request.rules_version, "data": base64.b64encode(pdf_bytes).decode('utf-8')}

@app.post("/api/rules", summary="Update business rules")
def update_rules_endpoint(new_rules: dict):
    global RULES_VERSION
    unsupported = set(new_rules) - set(EFFECTIVE_RULE_KEYS)
    if unsupported:
        raise HTTPException(status_code=400, detail=f"These settings are not currently effective and cannot be changed: {', '.join(sorted(unsupported))}")
    for key, value in new_rules.items():
        if not isinstance(value, str) or not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", value):
            raise HTTPException(status_code=422, detail=f"{key} must be a time in HH:MM format.")
    BUSINESS_RULES.update(new_rules)
    RULES_VERSION = effective_rules_version(BUSINESS_RULES)
    return {'success': True, 'rules': {key: BUSINESS_RULES[key] for key in EFFECTIVE_RULE_KEYS}, 'rules_version': RULES_VERSION}

@app.get("/api/rules", summary="Get current business rules")
def get_rules_endpoint():
    return {
        'rules': BUSINESS_RULES,
        'effective_rules': {key: BUSINESS_RULES[key] for key in EFFECTIVE_RULE_KEYS},
        'rules_version': effective_rules_version(BUSINESS_RULES),
        'unresolved': {
            'late_arrival_threshold': 'HR decision required: late minutes currently count from normal_start_time (08:00); 08:30 threshold is not applied.',
            'default_work_end_time': 'Not applied: missing checkout is left unknown and is not inferred from scheduled end.',
            'weekend_overtime_evaluated_separately': 'Weekend policy is unresolved; weekend overtime is withheld and flagged for review.'
        }
    }
