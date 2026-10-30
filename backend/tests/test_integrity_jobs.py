import base64
import asyncio
import csv
import io
import json
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook

import main


def event(kind, timestamp, date="2026-10-28", employee_id="E1"):
    try:
        parsed_timestamp = datetime.fromisoformat(timestamp) if timestamp else None
    except ValueError:
        parsed_timestamp = None
    return main._Event(employee_id, "Synthetic Employee", date, kind, parsed_timestamp, kind, timestamp)


def attendance_csv(employee_id, kind="IN", timestamp="2026-10-28 08:00:00"):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Employee ID", "Employee Name", "Event Type", "Date", "Timestamp"])
    writer.writerow([employee_id, "Synthetic Employee", kind, "2026-10-28", timestamp])
    if kind.upper() == "IN":
        writer.writerow([employee_id, "Synthetic Employee", "OUT", "2026-10-28", "2026-10-28 17:00:00"])
    return output.getvalue().encode()


class LocalResponse:
    def __init__(self, status_code, payload):
        self.status_code = status_code
        self.payload = payload
        self.text = json.dumps(payload, default=str)

    def json(self):
        return self.payload


class SyntheticUploadFile:
    def __init__(self, filename, content):
        self.filename = filename
        self.content = content

    async def read(self, size=-1):
        return self.content if size < 0 else self.content[:size]

    async def close(self):
        return None


class AttendanceIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.original_temp_dir = main.TEMP_DIR
        self.temp = tempfile.TemporaryDirectory(prefix="mm-synthetic-tests-")
        main.TEMP_DIR = Path(self.temp.name)
        self.original_rules = dict(main.BUSINESS_RULES)

    def tearDown(self):
        main.TEMP_DIR = self.original_temp_dir
        main.BUSINESS_RULES.clear()
        main.BUSINESS_RULES.update(self.original_rules)
        main.RULES_VERSION = main.effective_rules_version(main.BUSINESS_RULES)
        self.temp.cleanup()

    def post(self, path, *, json_body=None, data=None, files=None):
        """Call the actual route function directly; this environment's ASGI test transports hang."""
        try:
            if files:
                filename, content, _content_type = files["file"]
                upload = SyntheticUploadFile(filename, content)
                if path == "/api/upload":
                    payload = asyncio.run(main.upload_file(upload))
                else:
                    payload = asyncio.run(main.upload_roster(job_id=(data or {}).get("job_id", ""), file=upload))
            elif path == "/api/process":
                payload = main.process_attendance_endpoint(main.ProcessRequest(**(json_body or {})))
            elif path in ("/api/export/excel", "/api/export/pdf"):
                request = main.ExportRequest(**(json_body or {}))
                payload = main.export_excel_endpoint(request) if path.endswith("excel") else main.export_pdf_endpoint(request)
            elif path == "/api/rules":
                payload = main.update_rules_endpoint(json_body or {})
            else:
                raise AssertionError(f"Unsupported test route: {path}")
            return LocalResponse(200, payload)
        except HTTPException as error:
            return LocalResponse(error.status_code, {"detail": error.detail})
        except Exception as error:
            if error.__class__.__name__ == "ValidationError":
                return LocalResponse(422, {"detail": str(error)})
            raise

    def test_recognized_in_and_out_are_complete(self):
        day = main.process_attendance_events([
            event("check_in", "2026-10-28T08:00:00"),
            event("check_out", "2026-10-28T17:00:00"),
        ], main.BUSINESS_RULES)
        self.assertEqual(day["status"], "complete")
        self.assertFalse(day["flagged_for_review"])
        self.assertEqual(day["check_in"].strftime("%H:%M"), "08:00")
        self.assertEqual(day["check_out"].strftime("%H:%M"), "17:00")

    def test_roster_date_matching_attendance_does_not_duplicate_employee_day(self):
        days = main.process_employee_attendance([
            event("in", "2026-10-28T08:00:00", employee_id="SYN-100"),
            event("out", "2026-10-28T17:00:00", employee_id="SYN-100"),
        ], main.BUSINESS_RULES, {
            "SYN-100": {"name": "Synthetic Employee", "dates": ["2026-10-28"]},
        })
        self.assertEqual(len(days), 1)
        self.assertEqual(days[0]["status"], "complete")
        self.assertEqual(days[0]["employee_id"], "SYN-100")

    def test_unknown_event_types_never_fill_confirmed_punch_columns(self):
        day = main.process_attendance_events([
            event("Door held", "2026-10-28T08:00:00"),
            event("Access granted", "2026-10-28T17:00:00"),
        ], main.BUSINESS_RULES)
        self.assertIsNone(day["check_in"])
        self.assertIsNone(day["check_out"])
        self.assertNotEqual(day["status"], "complete")
        self.assertTrue(day["flagged_for_review"])
        self.assertEqual(day["event_evidence"][0]["event_type"], "Door held")

    def test_one_event_never_duplicates_one_timestamp_as_both_punches(self):
        day = main.process_attendance_events([event("in", "2026-10-28T08:00:00")], main.BUSINESS_RULES)
        self.assertIsNotNone(day["check_in"])
        self.assertIsNone(day["check_out"])
        self.assertNotEqual(day["status"], "complete")
        self.assertTrue(day["flagged_for_review"])

    def test_no_events_is_not_absence(self):
        day = main.process_attendance_events([], main.BUSINESS_RULES)
        self.assertEqual(day["status"], "no_events")
        self.assertIn("does not confirm absence", day["notes"])
        self.assertTrue(day["flagged_for_review"])

    def test_duplicate_timestamps_are_ambiguous_and_metrics_are_withheld(self):
        day = main.process_attendance_events([
            event("in", "2026-10-28T08:00:00"),
            event("out", "2026-10-28T08:00:00"),
        ], main.BUSINESS_RULES)
        self.assertEqual(day["status"], "ambiguous")
        self.assertTrue(day["flagged_for_review"])
        self.assertEqual(day["late_minutes"], 0)
        self.assertEqual(day["overtime_minutes"], 0)

    def test_missing_timestamp_is_preserved_and_flagged(self):
        day = main.process_attendance_events([
            event("in", "not-a-timestamp"), event("out", "2026-10-28T17:00:00"),
        ], main.BUSINESS_RULES)
        self.assertEqual(day["status"], "incomplete")
        self.assertTrue(day["flagged_for_review"])
        self.assertEqual(day["event_evidence"][0]["timestamp"], "not-a-timestamp")
        self.assertFalse(day["event_evidence"][0]["timestamp_valid"])

    def test_lateness_and_overtime_boundaries(self):
        def run(in_at, out_at):
            return main.process_attendance_events([
                event("in", f"2026-10-28T{in_at}:00"),
                event("out", f"2026-10-28T{out_at}:00"),
            ], main.BUSINESS_RULES)
        self.assertEqual(run("08:00", "19:30")["late_minutes"], 0)
        self.assertEqual(run("08:01", "19:30")["late_minutes"], 1)
        # The 08:30 HR policy is unresolved; the active metric remains minutes after 08:00.
        self.assertEqual(run("08:30", "19:30")["late_minutes"], 30)
        self.assertTrue(run("09:30", "19:30")["very_late"])
        self.assertEqual(run("08:00", "19:30")["overtime_minutes"], 0)
        self.assertEqual(run("08:00", "19:31")["overtime_minutes"], 1)

    def test_weekend_overtime_is_withheld_pending_policy(self):
        day = main.process_attendance_events([
            event("in", "2026-10-31T08:00:00", date="2026-10-31"),
            event("out", "2026-10-31T21:00:00", date="2026-10-31"),
        ], main.BUSINESS_RULES)
        self.assertEqual(day["overtime_minutes"], 0)
        self.assertTrue(day["flagged_for_review"])
        self.assertIn("Weekend overtime policy is unresolved", day["notes"])

    def upload(self, employee_id):
        response = self.post("/api/upload", files={"file": ("synthetic.csv", attendance_csv(employee_id), "text/csv")})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIs(response.json()["success"], True)
        return response.json()["job_id"]

    def test_jobs_and_rosters_are_isolated_and_exports_use_saved_job(self):
        job_a = self.upload("101")
        job_b = self.upload("202")
        roster_csv = b"id,name,work_date\n999,Synthetic Roster Person,2026-10-29\n"
        roster_response = self.post("/api/roster", data={"job_id": job_a}, files={"file": ("roster.csv", roster_csv, "text/csv")})
        self.assertEqual(roster_response.status_code, 200, roster_response.text)
        self.assertEqual(roster_response.json()["job_id"], job_a)

        result_a = self.post("/api/process", json_body={"job_id": job_a}).json()
        result_b = self.post("/api/process", json_body={"job_id": job_b}).json()
        self.assertEqual(result_a["job_id"], job_a)
        self.assertEqual(result_b["job_id"], job_b)
        self.assertIn("999", {row["employee_id"] for row in result_a["processed_days"]})
        self.assertNotIn("999", {row["employee_id"] for row in result_b["processed_days"]})
        self.assertIn("101", {row["employee_id"] for row in result_a["processed_days"]})
        self.assertIn("202", {row["employee_id"] for row in result_b["processed_days"]})

        payload = {"job_id": job_a, "rules_version": result_a["rules_version"]}
        changed = dict(main.BUSINESS_RULES)
        changed["normal_start_time"] = "08:05"
        update = self.post("/api/rules", json_body={key: changed[key] for key in main.EFFECTIVE_RULE_KEYS})
        self.assertEqual(update.status_code, 200, update.text)
        self.assertEqual(self.post("/api/export/excel", json_body=payload).status_code, 409)

        latest = self.post("/api/process", json_body={"job_id": job_a}).json()
        payload["rules_version"] = latest["rules_version"]
        exported = self.post("/api/export/excel", json_body=payload)
        self.assertEqual(exported.status_code, 200, exported.text)
        self.assertIs(exported.json()["success"], True)
        workbook = load_workbook(io.BytesIO(base64.b64decode(exported.json()["data"])))
        self.assertEqual(workbook.active.cell(1, 12).value, "Source event evidence")
        metadata_sheet = workbook["Report Metadata"]
        self.assertEqual(metadata_sheet.cell(1, 2).value, "synthetic.csv")
        self.assertEqual(metadata_sheet.cell(2, 2).value, job_a)
        self.assertEqual(metadata_sheet.cell(3, 2).value, latest["rules_version"])
        pdf = self.post("/api/export/pdf", json_body=payload)
        self.assertEqual(pdf.status_code, 200, pdf.text)
        self.assertIs(pdf.json()["success"], True)
        self.assertTrue(base64.b64decode(pdf.json()["data"]).startswith(b"%PDF"))

    def test_ambiguous_event_evidence_is_in_excel_export(self):
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Employee ID", "Employee Name", "Event Type", "Date", "Timestamp"])
        writer.writerow(["303", "Synthetic", "Unknown Punch", "2026-10-28", "2026-10-28 08:15:00"])
        response = self.post("/api/upload", files={"file": ("unknown.csv", output.getvalue().encode(), "text/csv")})
        job_id = response.json()["job_id"]
        result = self.post("/api/process", json_body={"job_id": job_id}).json()
        exported = self.post("/api/export/excel", json_body={"job_id": job_id, "rules_version": result["rules_version"]}).json()
        sheet = load_workbook(io.BytesIO(base64.b64decode(exported["data"]))).active
        self.assertEqual(sheet.cell(2, 4).value, "Unknown")
        self.assertEqual(sheet.cell(2, 5).value, "Unknown")
        self.assertIn("Unknown Punch", sheet.cell(2, 12).value)
        self.assertEqual(sheet.cell(2, 10).value, "incomplete")
        pdf = self.post("/api/export/pdf", json_body={"job_id": job_id, "rules_version": result["rules_version"]})
        self.assertEqual(pdf.status_code, 200, pdf.text)
        pdf_bytes = base64.b64decode(pdf.json()["data"])
        self.assertTrue(pdf_bytes.startswith(b"%PDF"))
        self.assertIn(b"Unknown Punch", pdf_bytes)

    def test_invalid_job_and_roster_uploads_are_rejected(self):
        self.assertEqual(self.post("/api/process", json_body={"job_id": "not-a-uuid"}).status_code, 400)
        self.assertEqual(self.post("/api/process", json_body={}).status_code, 422)
        self.assertEqual(self.post("/api/roster", data={"job_id": "not-a-uuid"}, files={"file": ("r.csv", b"id,name\n", "text/csv")}).status_code, 400)
        job_id = self.upload("404")
        unsupported = self.post("/api/roster", data={"job_id": job_id}, files={"file": ("roster.xls", b"synthetic", "application/octet-stream")})
        self.assertEqual(unsupported.status_code, 400)
        invalid_attendance = self.post("/api/upload", files={"file": ("attendance.txt", b"synthetic", "text/plain")})
        self.assertEqual(invalid_attendance.status_code, 400)
        unsupported_rule = self.post("/api/rules", json_body={"late_arrival_threshold": "08:30"})
        self.assertEqual(unsupported_rule.status_code, 400)
        malformed_rule = self.post("/api/rules", json_body={"normal_start_time": "25:99"})
        self.assertEqual(malformed_rule.status_code, 422)

    def test_oversized_roster_is_rejected(self):
        job_id = self.upload("505")
        content = b"x" * (10 * 1024 * 1024 + 1)
        response = self.post("/api/roster", data={"job_id": job_id}, files={"file": ("roster.csv", content, "text/csv")})
        self.assertEqual(response.status_code, 413)

    def test_concurrent_jobs_have_distinct_ids_and_results(self):
        def submit(employee_id):
            response = self.post("/api/upload", files={"file": ("parallel.csv", attendance_csv(employee_id), "text/csv")})
            return response.status_code, response.json()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(submit, ("606", "707")))
        self.assertTrue(all(status == 200 for status, _ in results))
        job_ids = [body["job_id"] for _, body in results]
        self.assertEqual(len(set(job_ids)), 2)
        for job_id in job_ids:
            processed = self.post("/api/process", json_body={"job_id": job_id})
            self.assertEqual(processed.status_code, 200, processed.text)
            row_ids = {day["employee_id"] for day in processed.json()["processed_days"]}
            self.assertEqual(len(row_ids), 1)


class AttendanceCsvValidationTests(unittest.TestCase):
    def assert_csv_rejected(self, content, message_part):
        with self.assertRaises(HTTPException) as raised:
            main.parse_ivms_csv(content.encode())
        self.assertEqual(raised.exception.status_code, 400)
        self.assertIn(message_part, raised.exception.detail)

    def test_unrelated_two_column_csv_is_rejected(self):
        self.assert_csv_rejected(
            "nonsense,not an attendance export\nalpha,beta\n",
            "headers must identify employee ID",
        )

    def test_empty_csv_is_rejected(self):
        self.assert_csv_rejected("", "Empty CSV file")

    def test_missing_required_headers_are_rejected(self):
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Department,Date\nSYN-1,Synthetic,IT,2026-10-28\n",
            "headers must identify employee ID, event type, and timestamp",
        )

    def test_missing_required_row_values_are_rejected(self):
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\nSYN-1,Synthetic,IN,2026-10-28,\n",
            "missing employee ID, event type, or timestamp",
        )
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\n,Synthetic,IN,2026-10-28,2026-10-28 08:00:00\n",
            "missing employee ID, event type, or timestamp",
        )

    def test_invalid_timestamp_and_date_are_rejected(self):
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\nSYN-1,Synthetic,IN,2026-10-28,not-a-time\n",
            "invalid or unsupported timestamp",
        )
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\nSYN-1,Synthetic,IN,not-a-date,2026-10-28 08:00:00\n",
            "invalid or unsupported date",
        )
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\nSYN-1,Synthetic,IN,2026-10-27,2026-10-28 08:00:00\n",
            "date does not match its timestamp date",
        )

    def test_valid_named_header_attendance_csv_is_accepted(self):
        parsed = main.parse_ivms_csv((
            "Employee ID,Employee Name,Event Type,Date,Timestamp\n"
            "SYN-1,Synthetic,IN,2026-10-28,2026-10-28 08:00:00\n"
            "SYN-1,Synthetic,OUT,2026-10-28,2026-10-28 17:00:00\n"
        ).encode())
        self.assertEqual(parsed["total_rows"], 2)
        self.assertEqual(len(main.extract_attendance_events(parsed)), 2)

    def test_legacy_four_column_positional_export_is_preserved(self):
        parsed = main.parse_ivms_csv((
            "Worker Ref,Person,Action,Time Recorded\n"
            "SYN-2,Synthetic,IN,2026-10-28 08:00:00\n"
            "SYN-2,Synthetic,OUT,2026-10-28 17:00:00\n"
        ).encode())
        self.assertEqual(parsed["column_map"]["employee_id"], 0)
        self.assertEqual(parsed["column_map"]["event_type"], 2)
        self.assertEqual(parsed["column_map"]["timestamp"], 3)
        self.assertEqual(len(main.extract_attendance_events(parsed)), 2)

    def test_invalid_row_rejects_whole_file_instead_of_partial_import(self):
        self.assert_csv_rejected(
            "Employee ID,Employee Name,Event Type,Date,Timestamp\n"
            "SYN-1,Synthetic,IN,2026-10-28,2026-10-28 08:00:00\n"
            "SYN-1,Synthetic,OUT,2026-10-28,bad-time\n",
            "row 3 has an invalid or unsupported timestamp",
        )

    def test_rejected_upload_does_not_leave_a_job_directory(self):
        original_temp_dir = main.TEMP_DIR
        with tempfile.TemporaryDirectory(prefix="mm-invalid-csv-") as temp_dir:
            main.TEMP_DIR = Path(temp_dir)
            try:
                upload = SyntheticUploadFile(
                    "invalid.csv",
                    b"nonsense,not an attendance export\nalpha,beta\n",
                )
                with self.assertRaises(HTTPException) as raised:
                    asyncio.run(main.upload_file(upload))
                self.assertEqual(raised.exception.status_code, 400)
                self.assertIn("headers must identify employee ID", raised.exception.detail)
                self.assertEqual(list(Path(temp_dir).iterdir()), [])
            finally:
                main.TEMP_DIR = original_temp_dir

